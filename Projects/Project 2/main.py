import openpyxl

inv_file = openpyxl.load_workbook("inventory.xlsx")
product_list = inv_file["Sheet1"]


# Task 1 -> Computing total product supplied by each supplier
product_per_supplier = {}

for product_row in range(2, product_list.max_row+1):
    supplier_name = product_list.cell(product_row, 4).value

    if supplier_name in product_per_supplier:
        product_per_supplier[supplier_name] = product_per_supplier[supplier_name] + 1
    else:
        print("Adding a new supplier:", supplier_name)
        product_per_supplier[supplier_name] = 1

print(product_per_supplier)
print("------------X--------------\n")

# Task 2 -> Computing total inventory value per supplier
inventory_value_per_supplier = {}
for product_row in range(2, product_list.max_row+1):
    supplier_name = product_list.cell(product_row, 4).value
    inventory_quantity = product_list.cell(product_row, 2).value
    inventory_value = product_list.cell(product_row, 3).value

    if supplier_name in inventory_value_per_supplier:
        inventory_value_per_supplier[supplier_name] += inventory_value * inventory_quantity
    else:
        inventory_value_per_supplier[supplier_name] =  inventory_value * inventory_quantity

print(inventory_value_per_supplier)
print("------------X--------------\n")

# Task 3 -> Listing products with inventory quantity less than 10
product_with_inventory_LessThan10 = {}
for product_row in range(2, product_list.max_row+1):
    supplier_name = product_list.cell(product_row, 4).value
    inv_quantity = product_list.cell(product_row, 2).value
    product_no = product_list.cell(product_row, 1).value

    if inv_quantity < 10:
        product_with_inventory_LessThan10[int(product_no)] = int(inv_quantity)

print(product_with_inventory_LessThan10)
print("------------X--------------\n")

# Task 4 -> Write total inventory value for a particular product
for product_row in range(2, product_list.max_row+1):
    inventory_price = product_list.cell(product_row, 5)
    inventory_quantity = product_list.cell(product_row, 2).value
    inventory_value = product_list.cell(product_row, 3).value

    inventory_price.value = inventory_quantity * inventory_value

inv_file.save("inventory_with_total_value.xlsx")